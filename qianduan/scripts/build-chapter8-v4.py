from __future__ import annotations

import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "新第8章-随钻围岩状态智能反演三维云图平台构建_v3.docx"
OUTPUT = ROOT / "新第8章-随钻围岩状态智能反演三维云图平台构建_v4.docx"


def set_run_font(run, chinese: str, size: float, *, bold: bool | None = None) -> None:
    run.font.name = chinese
    run._element.rPr.rFonts.set(qn("w:eastAsia"), chinese)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def set_keep_together(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepLines")) is None:
        p_pr.append(OxmlElement("w:keepLines"))


def add_body(doc: Document, text: str, *, bold_lead: str | None = None):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles["Normal"]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    set_keep_together(paragraph)

    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, "黑体", 12, bold=True)
        rest = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(rest, "宋体", 12)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, "宋体", 12)
    return paragraph


def add_heading(doc: Document, text: str, level: int):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(8 if level == 2 else 5)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.0
    set_keep_with_next(paragraph)
    size = {1: 16, 2: 15, 3: 13.5}[level]
    for run in paragraph.runs:
        set_run_font(run, "黑体", size, bold=True)
    return paragraph


def add_equation(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    set_keep_together(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, "Cambria Math", 11.5)
    return paragraph


def add_figure(doc: Document, image_path: Path, caption: str, width: float = 6.35):
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.first_line_indent = Pt(0)
    picture.paragraph_format.space_before = Pt(4)
    picture.paragraph_format.space_after = Pt(2)
    set_keep_with_next(picture)
    picture.add_run().add_picture(str(image_path), width=Inches(width))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Pt(0)
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(5)
    set_keep_together(cap)
    run = cap.add_run(caption)
    set_run_font(run, "宋体", 10.5)
    return picture, cap


def set_cell_text(cell, text: str, *, bold: bool = False, size: float = 9.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(text)
    set_run_font(run, "黑体" if bold else "宋体", size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge, val, size in (
        ("top", "single", "12"),
        ("left", "nil", "0"),
        ("bottom", "single", "12"),
        ("right", "nil", "0"),
        ("insideH", "single", "4"),
        ("insideV", "nil", "0"),
    ):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), val)
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), "808080")
        borders.append(element)
    tbl_pr.append(borders)


def add_model_table(doc: Document):
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(2)
    set_keep_with_next(caption)
    run = caption.add_run("表 8-1  三种模型的元数据评价指标")
    set_run_font(run, "宋体", 10.5)

    headers = ["模型", "损伤准确率", "应力准确率", "状态准确率", "宏平均 F1"]
    rows = [
        ["V1 多尺度极端随机树", "73.03%", "85.76%", "68.40%", "68.08%"],
        ["V2 CNN-BiLSTM", "70.33%", "80.47%", "69.42%", "66.67%"],
        ["V3 物理融合 CNN-BiLSTM", "71.76%", "83.12%", "69.90%", "68.44%"],
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(2.15), Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.0)]
    for index, header in enumerate(headers):
        table.columns[index].width = widths[index]
        set_cell_text(table.rows[0].cells[index], header, bold=True)
        table.rows[0].cells[index].width = widths[index]
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "E8EEF2")
        table.rows[0].cells[index]._tc.get_or_add_tcPr().append(shade)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            set_cell_text(cells[index], value)
            cells[index].width = widths[index]
    set_table_borders(table)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note.paragraph_format.first_line_indent = Pt(0)
    note.paragraph_format.space_before = Pt(2)
    note.paragraph_format.space_after = Pt(4)
    run = note.add_run("注：指标来自数据文件随附元数据，仅用于同源数据条件下的模型比较。")
    set_run_font(run, "宋体", 9.5)


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def normalize_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)

    for level, size in ((1, 16), (2, 15), (3, 13.5)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)


def build() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    doc = Document(SOURCE)
    clear_body(doc)
    normalize_styles(doc)

    props = doc.core_properties
    props.title = "第八章 随钻智控平台构建与三维时序表征"
    props.subject = "随钻参数、围岩状态反演、三维时序体云与智能卸压决策接口"
    props.keywords = "随钻智控; 围岩状态; 三维时序体云; 钻孔反演"

    with tempfile.TemporaryDirectory(prefix="chapter8_v4_") as temp_dir:
        temp = Path(temp_dir)
        with zipfile.ZipFile(SOURCE) as archive:
            image_names = [
                "word/media/image1.png",
                "word/media/image2.png",
                "word/media/image3.png",
                "word/media/image4.png",
            ]
            image_paths = []
            for name in image_names:
                target = temp / Path(name).name
                target.write_bytes(archive.read(name))
                image_paths.append(target)

        add_heading(
            doc,
            "8 “随钻智控”——基于随钻参数的深部强扰动围岩状态精细感知与智能卸压决策平台",
            1,
        )
        add_body(
            doc,
            "本章在前述围岩状态识别与预测研究的基础上，面向随钻参数组织、反演结果空间表达和交互分析需求，构建“随钻智控”平台。平台名称中的智能卸压决策体现其目标方向，但现有数据尚未包含明确的卸压措施参数和卸压效果标签。因此，本章重点讨论围岩状态感知、多模型结果集成与三维可视分析，卸压决策作为后续扩展接口，不将变应力工况直接解释为卸压效果。",
        )

        add_heading(doc, "8.1 研究目标与总体方案", 2)
        add_heading(doc, "8.1.1 研究对象与表达边界", 3)
        add_body(
            doc,
            "研究对象为环向布置钻孔的随钻参数及其状态反演结果。每组数据对应一个钻孔，共形成11个环向位置。钻进深度为沿巷道断面径向向围岩内部延伸的距离，范围为0～125 cm，按0.5 cm间隔形成251个深度帧。平台以θ表示钻孔环向位置，以r表示径向钻进深度，以Xₜ表示界面中的演进展开轴。Xₜ并非巷道纵向的实测空间坐标。",
        )
        add_body(
            doc,
            "源数据以深度帧组织，未提供可校准的绝对时间戳。本文所称时序是钻进过程的先后次序，用于表达不同钻进深度下状态场的连续演进。三维云图服务于结果解释和对比，不等同于经过现场标定的连续物理场，也不能替代独立的地质反演与力学验证。",
        )

        add_heading(doc, "8.1.2 系统架构与数据流程", 3)
        add_body(
            doc,
            "平台采用数据层、模型层、可视化层以及交互与决策层的四层架构。数据从原始随钻记录进入清洗和拟合流程，经统一深度网格组织后接入三种状态反演模型，随后映射为环向钻孔、径向深度和演进次序共同约束的三维表达，最终通过模型切换、场量切换和切面锁定完成交互分析。",
        )
        add_body(
            doc,
            "数据层。负责读取 processed_dataset_Vtest4.zip 与 拟合.zip，保留扭矩、推力、损伤、应力和样本来源等字段，并建立钻孔编号、环向角度与深度帧之间的对应关系。",
            bold_lead="数据层。",
        )
        add_body(
            doc,
            "模型层。集成 V1 多尺度极端随机树、V2 CNN-BiLSTM 和 V3 物理融合 CNN-BiLSTM，将损伤、应力、组合状态及置信度统一封装为可比较输出。",
            bold_lead="模型层。",
        )
        add_body(
            doc,
            "可视化层。以17个半透明切面、连续包络体和11条钻孔演进轨迹构成时序体云，并提供应力、损伤和综合误差三类场量表达。",
            bold_lead="可视化层。",
        )
        add_body(
            doc,
            "交互与决策层。支持演进播放、暂停锁定、深度拖拽、切面选择、视角控制和模型对比。当前版本提供状态研判所需的信息组织能力，尚未输出自动卸压方案。",
            bold_lead="交互与决策层。",
        )
        add_figure(doc, image_paths[0], "图 8-1  平台四层技术架构与数据流", width=6.4)

        add_heading(doc, "8.2 数据组织与反演结果集成", 2)
        add_heading(doc, "8.2.1 数据标准化与钻孔映射", 3)
        add_body(
            doc,
            "两组数据共包含82555条原始记录。预处理流程依次采用 Hampel 异常值识别、Savitzky–Golay 平滑和 PCHIP 分段三次 Hermite 插值，并将各钻孔重采样到统一深度网格。元数据记录了33组拟合过程，平均粗糙度下降率为72.85%。该指标反映曲线平滑程度的变化，不代表预测精度提高。",
        )
        add_body(
            doc,
            "11个钻孔按照环向角度排列，其中5个为主工况钻孔，5个为重复工况钻孔，另有1个变应力复合钻孔。各钻孔均包含251个深度帧。重复工况用于增强相同围压条件下的空间覆盖，复合钻孔 S99 的应力水平在10～40 MPa之间变化。由于数据中没有卸压钻孔标识、卸压量或卸压前后对照，S99仅作为特殊工况单独标注。",
        )

        add_heading(doc, "8.2.2 模型输出与评价原则", 3)
        add_body(
            doc,
            "平台将三种模型的损伤预测、应力预测和组合状态映射到同一数据结构。损伤值按0～80归一化，应力值按0～40 MPa归一化，综合误差取损伤绝对误差与应力绝对误差归一化后的均值。统一量纲有利于颜色映射和模型切换，但不会改变原始预测结果。",
        )
        add_model_table(doc)
        add_body(
            doc,
            "从随附指标看，V1的损伤准确率和应力准确率最高，V3的状态准确率与宏平均 F1 略高于其余模型。不同指标指向不同优势，因此平台保留模型切换功能，不以单一准确率直接判定最优模型。由于现有材料未给出完整的数据划分、重复试验和置信区间，表中数值不宜外推为现场泛化性能。",
        )

        add_heading(doc, "8.3 三维时序体云构建方法", 2)
        add_heading(doc, "8.3.1 场量映射与切面组织", 3)
        add_body(
            doc,
            "每个切面对应一个钻进演进位置，切面内部由环向角度与径向深度共同确定场量。对任意环向位置，平台根据该位置与11个钻孔之间的最短圆周角距离计算高斯型权重，再对各钻孔在同一径向深度处的归一化结果进行加权平均。其基本形式为：",
        )
        add_equation(doc, "F(θ,r,t) = Σ wᵢ(θ) fᵢ(r,t) / Σ wᵢ(θ)")
        add_body(
            doc,
            "权重扩散尺度取平均钻孔角间距的0.78倍，并加入极小基准权重以避免局部空值。该方法保证环向颜色连续，适合展示离散钻孔之间的趋势联系，但属于可视化插值，不包含岩体本构关系、边界条件或空间协方差模型。",
        )
        add_figure(doc, image_paths[1], "图 8-2  多切面连续时序体云主界面", width=6.35)

        add_heading(doc, "8.3.2 连续包络与局部切面分析", 3)
        add_body(
            doc,
            "平台沿演进展开轴设置17个半透明环形切面，相邻切面的代表深度间隔约为7.81 cm。切面之间通过连续包络体连接，并以11条低透明度轨迹保持钻孔位置在各切面之间的对应关系。已演进切面、当前切面和锁定分析面采用不同透明度与强调色，从而在整体结构中保留明确的时间顺序。",
        )
        add_body(
            doc,
            "单切面模式隐藏非目标切面，突出当前深度的环向分布和径向梯度。用户可点击切面或时间轨道选定分析对象，并结合模型、场量和钻孔信息进行局部比较。连续表面主要用于建立视觉联系，钻孔稀疏区的细节仍受插值假设影响，不能据此识别小于钻孔间距的异常结构。",
        )
        add_figure(doc, image_paths[2], "图 8-3  单切面锁定与局部分析界面", width=6.35)

        add_heading(doc, "8.4 交互分析与系统实现", 2)
        add_heading(doc, "8.4.1 钻进演进与暂停机制", 3)
        add_body(
            doc,
            "演进控制将0～125 cm径向钻深映射为0～100%的进度。自动播放按慢速、标准和快速三档推进；暂停操作同时锁定当前分析面；拖动深度滑块或选择任一切面也会自动暂停，防止分析对象在读数过程中继续变化。用户恢复播放后，分析面重新跟随当前演进位置。",
        )
        add_body(
            doc,
            "场量切换用于比较应力、损伤与综合误差，模型切换用于观察三种反演结果的差异。钻孔悬停和选择功能提供数据来源、样本编号与预测状态的追溯入口。由此形成整体演进观察、目标切面锁定、局部钻孔核查和跨模型复核的分析路径。",
        )
        add_figure(doc, image_paths[3], "图 8-4  损伤场量与多模型评价界面", width=6.35)

        add_heading(doc, "8.4.2 前端实现与数据更新", 3)
        add_body(
            doc,
            "平台采用 Vue 3 构建界面，Pinia 管理模型、钻孔和演进状态，Three.js 负责三维几何、材质、标注和交互拾取。数据处理脚本将两个压缩数据源转换为结构化 JSON，前端按需读取并在浏览器内完成颜色更新、切面显隐和轨迹重建。组件化设计使数据层、状态层与三维渲染层保持相对独立，便于后续接入实时采集接口。",
        )
        add_body(
            doc,
            "当前实现已通过项目生产构建与浏览器页面检查，但尚未开展统一硬件条件下的帧率、内存占用和响应时延基准测试。因此，本章只确认功能链路可以运行，不对实时性能给出未经测量的量化结论。",
        )

        add_heading(doc, "8.5 验证结果、局限性与本章小结", 2)
        add_heading(doc, "8.5.1 功能验证结果", 3)
        add_body(
            doc,
            "平台能够读取两组数据源，完成11个钻孔的环向映射和251个深度帧的统一索引，并在三种模型之间切换应力、损伤和综合误差场。多切面播放、暂停锁定、深度拖拽、单切面分析和钻孔选择均形成闭环交互。上述结果验证了数据到界面的功能一致性，不等同于对插值场真实性和模型现场适用性的独立验证。",
        )

        add_heading(doc, "8.5.2 研究局限与后续工作", 3)
        add_body(
            doc,
            "现阶段主要存在三方面限制。第一，时序轴由深度帧构造，缺少绝对时间戳和巷道纵向实测坐标。第二，环向连续场由11个离散钻孔插值得到，尚未量化空间不确定性。第三，数据中没有明确的卸压措施与效果标签，平台只能辅助状态研判，不能自动生成经验证的卸压参数。",
        )
        add_body(
            doc,
            "后续研究应接入带时间戳和真实孔位坐标的现场数据，引入交叉验证、重复试验及置信区间，并采用留孔验证或地质统计方法评估空间插值误差。在此基础上，可进一步采集卸压孔径、孔深、间距、施工时刻以及卸压前后应力变化，建立状态识别、方案推荐和效果反馈相衔接的决策闭环。",
        )

        add_heading(doc, "8.5.3 本章小结", 3)
        add_body(
            doc,
            "本章构建了面向随钻参数的围岩状态三维可视分析平台，明确了环向位置、径向钻深与演进次序的坐标含义，完成了11个钻孔、251个深度帧和三种模型结果的统一组织。平台通过17个半透明切面、连续包络体和钻孔轨迹增强时序联系，并以暂停锁定和单切面选择支持局部分析。同时，本章对插值、模型评价和卸压数据的适用边界进行了限定，为后续现场验证和智能卸压决策模型接入提供了可扩展基础。",
        )

        doc.save(OUTPUT)


if __name__ == "__main__":
    try:
        build()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
    print(OUTPUT)
